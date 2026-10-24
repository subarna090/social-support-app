"""Assets and liabilities spreadsheet (research R7). Totals are values, not formulas."""

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from social_support.synthetic.documents.common import MARKING, RenderedDocument

MONEY_FORMAT = "#,##0.00"


def _sheet(workbook: Workbook, title: str, header: list[str], rows: list[list], first: bool):
    sheet = workbook.active if first else workbook.create_sheet()
    sheet.title = title
    sheet["A1"] = MARKING
    sheet["A1"].font = Font(bold=True, color="B00000")
    sheet.append([])
    sheet.append(header)
    for cell in sheet[3]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for column in sheet.iter_cols(min_row=4):
        for cell in column:
            if isinstance(cell.value, float):
                cell.number_format = MONEY_FORMAT
    for letter, width in zip("ABCD", (22, 32, 26, 24), strict=True):
        sheet.column_dimensions[letter].width = width
    return sheet


def render_assets_liabilities(shown: dict, path: Path, as_of: date) -> RenderedDocument:
    workbook = Workbook()
    _sheet(
        workbook,
        "Assets",
        ["Type", "Description", "Value (AED)"],
        [
            [a["kind"].title(), a["description"], float(Decimal(a["value_aed"]))]
            for a in shown["assets"]
        ],
        first=True,
    )
    _sheet(
        workbook,
        "Liabilities",
        ["Type", "Lender", "Outstanding balance (AED)", "Monthly payment (AED)"],
        [
            [
                li["kind"].replace("_", " ").title(),
                li["lender"],
                float(Decimal(li["outstanding_aed"])),
                float(Decimal(li["monthly_payment_aed"])),
            ]
            for li in shown["liabilities"]
        ],
        first=False,
    )
    _sheet(
        workbook,
        "Summary",
        ["Item", "Amount (AED)"],
        [
            ["Total assets (AED)", float(Decimal(shown["total_assets_aed"]))],
            ["Total liabilities (AED)", float(Decimal(shown["total_liabilities_aed"]))],
            ["Net wealth (AED)", float(Decimal(shown["net_wealth_aed"]))],
        ],
        first=False,
    )
    stamp = datetime(as_of.year, as_of.month, as_of.day)
    workbook.properties.created = stamp
    workbook.properties.modified = stamp
    workbook.properties.creator = "social-support-data"
    workbook.save(path)
    return RenderedDocument(path, shown)
