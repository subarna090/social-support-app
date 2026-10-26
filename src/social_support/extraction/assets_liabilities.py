"""Assets and liabilities spreadsheet: the totals on its Summary sheet (research R1)."""

import io
import zipfile

from openpyxl import load_workbook

from social_support.domain.documents import DocumentType, ExtractedDocument
from social_support.extraction.common import money, result
from social_support.extraction.pdf_text import ExtractionError

ROWS = {
    "Total assets (AED)": "total_assets_aed",
    "Total liabilities (AED)": "total_liabilities_aed",
    "Net wealth (AED)": "net_wealth_aed",
}


def parse(data: bytes) -> ExtractedDocument:
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ExtractionError("the spreadsheet could not be opened") from exc
    fields: dict[str, str] = {}
    if "Summary" in book.sheetnames:
        for row in book["Summary"].iter_rows(values_only=True):
            if len(row) >= 2 and row[0] in ROWS and isinstance(row[1], int | float):
                fields[ROWS[row[0]]] = money(row[1])
    book.close()
    return result(DocumentType.ASSETS_LIABILITIES, fields)
