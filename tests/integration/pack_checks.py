"""Reusable checks: does each generated document agree with its pack's answer key? (T014, T027)

Answer keys hold canonical values ("6000.00", YYYY-MM-DD). Documents display them through the
renderers' own helpers, so values are formatted with ``money``/``month_label`` before searching.
"""

import json
import re
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from social_support.synthetic.documents.common import MARKING, money, month_label

PDFS = ("bank_statement", "resume", "credit_report")


def load(folder: Path) -> tuple[dict, dict]:
    form = json.loads((folder / "form.json").read_text(encoding="utf-8"))
    key = json.loads((folder / "answer_key.json").read_text(encoding="utf-8"))
    return form, key


def pdf_pages(path: Path) -> list[str]:
    return [re.sub(r"\s+", " ", page.extract_text() or "") for page in PdfReader(path).pages]


def pdf_text(path: Path) -> str:
    return " ".join(pdf_pages(path))


def one_line(address: dict) -> str:
    return f"{address['unit']}, {address['street']}, {address['district']}, {address['emirate']}"


def shown_values_appear(folder: Path, key: dict) -> list[str]:
    """Each document's displayed values (answer_key.documents) really are in the file."""
    docs, problems = key["documents"], []

    def expect(text: str, value: str, where: str) -> None:
        if value not in text:
            problems.append(f"{where}: {value!r} not found")

    bank = pdf_text(folder / "bank_statement.pdf")
    b = docs["bank_statement"]
    expect(bank, b["account_holder"], "bank_statement")
    expect(bank, b["address"], "bank_statement")
    expect(bank, money(Decimal(b["monthly_income_credits_aed"])), "bank_statement")
    for credit in b["income_credits"]:
        expect(bank, money(Decimal(credit["amount_aed"])), "bank_statement")

    resume = pdf_text(folder / "resume.pdf")
    r = docs["resume"]
    expect(resume, r["full_name"], "resume")
    for job in r["jobs"]:
        expect(resume, job["employer"], "resume")
        expect(resume, month_label(job["start"]), "resume")

    credit = pdf_text(folder / "credit_report.pdf")
    c = docs["credit_report"]
    for field in ("full_name", "emirates_id", "date_of_birth", "address"):
        expect(credit, c[field], "credit_report")
    expect(credit, str(c["credit_score"]), "credit_report")
    for member in c["household_members"]:
        expect(credit, member["name"], "credit_report")

    summary = load_workbook(folder / "assets_liabilities.xlsx")["Summary"]
    cells = {row[0]: row[1] for row in summary.iter_rows(min_row=3, values_only=True) if row[0]}
    a = docs["assets_liabilities"]
    for label, field in (
        ("Total assets (AED)", "total_assets_aed"),
        ("Total liabilities (AED)", "total_liabilities_aed"),
        ("Net wealth (AED)", "net_wealth_aed"),
    ):
        if Decimal(str(cells.get(label))).quantize(Decimal("0.01")) != Decimal(a[field]):
            problems.append(
                f"assets_liabilities: {label} is {cells.get(label)}, expected {a[field]}"
            )
    return problems


def marking_problems(folder: Path) -> list[str]:
    problems = []
    for name in PDFS:
        for number, page in enumerate(pdf_pages(folder / f"{name}.pdf"), 1):
            if MARKING not in page:
                problems.append(f"{name}.pdf page {number} has no synthetic marking")
    workbook = load_workbook(folder / "assets_liabilities.xlsx")
    for sheet in workbook.worksheets:
        if sheet["A1"].value != MARKING:
            problems.append(f"assets_liabilities.xlsx sheet {sheet.title} has no marking in A1")
    return problems


def differences(key: dict, form: dict) -> set[tuple[str, str]]:
    """(field, document) pairs where a document disagrees with the true profile."""
    profile, docs, facts = key["profile"], key["documents"], key["facts"]
    found: set[tuple[str, str]] = set()

    def compare(field: str, truth, shown: dict[str, object]) -> None:
        for document, value in shown.items():
            if value != truth:
                found.add((field, document))

    card, credit, bank, resume = (
        docs["emirates_id"],
        docs["credit_report"],
        docs["bank_statement"],
        docs["resume"],
    )
    compare(
        "full_name",
        profile["full_name"],
        {
            "form": form["full_name"],
            "emirates_id": card["full_name"],
            "credit_report": credit["full_name"],
            "bank_statement": bank["account_holder"],
            "resume": resume["full_name"],
        },
    )
    compare(
        "emirates_id",
        profile["emirates_id"],
        {
            "form": form["emirates_id"],
            "emirates_id": card["emirates_id"],
            "credit_report": credit["emirates_id"],
        },
    )
    compare(
        "date_of_birth",
        profile["date_of_birth"],
        {
            "form": form["date_of_birth"],
            "emirates_id": card["date_of_birth"],
            "credit_report": credit["date_of_birth"],
        },
    )
    address = one_line(profile["address"])
    compare(
        "address",
        address,
        {
            "form": one_line(form["address"]),
            "credit_report": credit["address"],
            "bank_statement": bank["address"],
            "resume": resume["address"],
        },
    )
    current = next((j["employer"] for j in profile["jobs"] if j["end"] is None), None)
    compare("employer", current, {"form": form["employer"], "resume": resume["current_employer"]})
    compare(
        "net_monthly_income_aed",
        facts["net_monthly_income_aed"],
        {
            "form": form["net_monthly_income_aed"],
            "bank_statement": bank["monthly_income_credits_aed"],
        },
    )
    members = sorted((m["name"], m["relationship"], m["age"]) for m in profile["members"])
    compare(
        "family_members",
        members,
        {
            "form": sorted(
                (m["name"], m["relationship"], m["age"]) for m in form["family_members"]
            ),
            "credit_report": sorted(
                (m["name"], m["relationship"], m["age"]) for m in credit["household_members"]
            ),
        },
    )
    compare(
        "family_size",
        1 + len(profile["members"]),
        {
            "form": form["family_size"],
            "form_members": 1 + len(form["family_members"]),
        },
    )
    return found
