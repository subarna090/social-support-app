"""T015: same count, seed and settings give identical content (FR-006, SC-002, research R5)."""

import json
import re

import pytest
from openpyxl import load_workbook

from social_support.synthetic.run import GenerationRequest, generate

BYTE_IDENTICAL = (
    "form.json",
    "answer_key.json",
    "bank_statement.pdf",
    "resume.pdf",
    "credit_report.pdf",
    "emirates_id.png",
)


@pytest.fixture(scope="module")
def runs(tmp_path_factory):
    first = generate(GenerationRequest(count=5, seed=42, output_root=tmp_path_factory.mktemp("a")))
    second = generate(GenerationRequest(count=5, seed=42, output_root=tmp_path_factory.mktemp("b")))
    other = generate(GenerationRequest(count=5, seed=43, output_root=tmp_path_factory.mktemp("c")))
    return first, second, other


def test_run_id_is_deterministic(runs):
    first, second, _ = runs
    assert re.fullmatch(r"s42-n5-[0-9a-f]{6}", first.run_id)
    assert first.run_id == second.run_id


def test_summary_is_byte_identical(runs):
    first, second, _ = runs
    assert (first.folder / "summary.csv").read_bytes() == (
        second.folder / "summary.csv"
    ).read_bytes()


def test_files_are_byte_identical(runs):
    first, second, _ = runs
    for index in range(1, 6):
        a = first.folder / "packs" / f"pack-{index:04d}"
        b = second.folder / "packs" / f"pack-{index:04d}"
        for name in BYTE_IDENTICAL:
            assert (a / name).read_bytes() == (b / name).read_bytes(), f"{a.name}/{name}"


def test_spreadsheets_have_identical_cells(runs):
    first, second, _ = runs
    for index in range(1, 6):
        name = f"packs/pack-{index:04d}/assets_liabilities.xlsx"
        assert _cells(first.folder / name) == _cells(second.folder / name)


def test_run_record_differs_only_in_generated_at(runs):
    first, second, _ = runs
    a, b = (
        json.loads((r.folder / "run.json").read_text(encoding="utf-8")) for r in (first, second)
    )
    a.pop("generated_at"), b.pop("generated_at")
    assert a == b


def test_a_different_seed_gives_different_packs(runs):
    first, _, other = runs
    a = (first.folder / "packs/pack-0001/answer_key.json").read_text(encoding="utf-8")
    b = (other.folder / "packs/pack-0001/answer_key.json").read_text(encoding="utf-8")
    assert a != b


def _cells(path):
    workbook = load_workbook(path)
    return {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets}
