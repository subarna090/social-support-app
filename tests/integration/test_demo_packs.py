"""T012: the ``demo`` command against specs/004-poc-demo/contracts/cli.md."""

import json

from openpyxl import load_workbook

from social_support.synthetic.cli import main

PACK_FILES = {
    "bank_statement.pdf",
    "emirates_id.png",
    "resume.pdf",
    "assets_liabilities.xlsx",
    "credit_report.pdf",
    "form.json",
    "answer_key.json",
}
EXPECTED = {
    "approve": "APPROVE",
    "soft_decline": "SOFT_DECLINE",
    "review": "NEEDS_HUMAN_REVIEW",
}


def _content(folder):
    """Bytes per file; spreadsheets by cell values, because xlsx zips carry write timestamps."""
    result = {}
    for path in sorted(p for p in folder.rglob("*") if p.is_file()):
        name = path.relative_to(folder).as_posix()
        if path.suffix == ".xlsx":
            book = load_workbook(path)
            result[name] = [list(sheet.iter_rows(values_only=True)) for sheet in book]
        else:
            result[name] = path.read_bytes()
    return result


def test_demo_writes_three_packs_with_the_intended_outcomes(tmp_path, capsys):
    assert main(["demo", "--out", str(tmp_path)]) == 0
    out = capsys.readouterr().out
    assert {p.name for p in tmp_path.iterdir()} == set(EXPECTED)
    for name, outcome in EXPECTED.items():
        folder = tmp_path / name
        assert {p.name for p in folder.iterdir()} == PACK_FILES
        key = json.loads((folder / "answer_key.json").read_text(encoding="utf-8"))
        assert key["expected_routing"] == outcome
        form = json.loads((folder / "form.json").read_text(encoding="utf-8"))
        for secret in (form["full_name"], form["emirates_id"], form["net_monthly_income_aed"]):
            assert secret not in out  # console stays free of personal data


def test_review_pack_under_reports_income_on_the_form(tmp_path):
    main(["demo", "--out", str(tmp_path)])
    key = json.loads((tmp_path / "review" / "answer_key.json").read_text(encoding="utf-8"))
    assert key["inconsistency"]["type"] == "INCOME_MISMATCH"
    assert key["documents"]["form"]["net_monthly_income_aed"] == "7000.00"
    assert key["documents"]["bank_statement"]["monthly_income_credits_aed"] == "9000.00"


def test_demo_is_reproducible_and_needs_force_to_overwrite(tmp_path, capsys):
    first, second = tmp_path / "a", tmp_path / "b"
    assert main(["demo", "--out", str(first)]) == 0
    assert main(["demo", "--out", str(second)]) == 0
    assert _content(first) == _content(second)

    capsys.readouterr()
    assert main(["demo", "--out", str(first)]) == 2
    assert "--force" in capsys.readouterr().err
    assert main(["demo", "--out", str(first), "--force"]) == 0
    assert _content(first) == _content(second)
